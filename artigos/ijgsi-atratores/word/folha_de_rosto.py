"""Folha de rosto da IJGSI (arquivo separado do manuscrito anônimo).

Uso: python word/folha_de_rosto.py saida.docx
Campos entre colchetes precisam ser preenchidos pelos autores.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

tex = Path(__file__).resolve().parent.parent.joinpath("main.tex").read_text(encoding="utf8")
titulo = " ".join(re.sub(r"\\textbf\{|\\textit\{|\}", "", re.search(r"\\title\{(.*?)\}\}\}", tex, re.S).group(1)).split())

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, m, Cm(2.54))
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(6)


def par(texto="", negrito=False, italico=False, caixa_alta=False, alin=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = alin
    if texto:
        r = p.add_run(texto)
        r.bold, r.italic = negrito, italico
        r.font.all_caps = caixa_alta
    return p


def campo(rotulo, valor):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(rotulo).bold = True
    p.add_run(valor)


par(titulo, negrito=True, italico=True, caixa_alta=True)
par("Title page", negrito=True, alin=WD_ALIGN_PARAGRAPH.LEFT)

par("Authors", negrito=True, alin=WD_ALIGN_PARAGRAPH.LEFT)
for nome, email in (("Caio Silva Azeredo", "caiosazeredo@cos.ufrj.br"), ("Geraldo Xexéo", "[e-mail]")):
    p = doc.add_paragraph()
    p.add_run(nome).bold = True
    p.add_run("\nSystems Engineering and Computer Science Program (PESC/COPPE), Federal University of Rio de Janeiro (UFRJ), Rio de Janeiro, Brazil")
    p.add_run(f"\nE-mail: {email} · ORCID: [https://orcid.org/0000-0000-0000-0000]")

par("Short biographies (100–150 words each)", negrito=True, alin=WD_ALIGN_PARAGRAPH.LEFT)
campo("Caio Silva Azeredo. ", "[100–150 words: current position, research interests, relevant experience.]")
campo("Geraldo Xexéo. ", "[100–150 words: current position, research interests, relevant experience.]")

par("Acknowledgements", negrito=True, alin=WD_ALIGN_PARAGRAPH.LEFT)
par("[Optional: funding agencies and grant numbers, e.g. CAPES/CNPq/FAPERJ; people who helped.]")

par("Conflict of interest statement", negrito=True, alin=WD_ALIGN_PARAGRAPH.LEFT)
par("The authors declare no conflict of interest.")

par("Use of generative AI", negrito=True, alin=WD_ALIGN_PARAGRAPH.LEFT)
par("Language models are the object of study and were used to generate the simulated sessions and as judges in the "
    "validation, as described in the manuscript. [Declare here any use of AI tools in writing, analysis or code, "
    "according to the journal's policy and to what the authors actually did.]")

par("Authorship statement", negrito=True, alin=WD_ALIGN_PARAGRAPH.LEFT)
par("Caio Silva Azeredo: Conceptualization, Methodology, Software, Investigation, Formal analysis, Writing - Original draft. "
    "Geraldo Xexéo: Conceptualization, Supervision, Writing - Reviewing and Editing. [Confirm the roles.]")

cp = doc.core_properties
cp.author = cp.last_modified_by = cp.title = cp.comments = ""
doc.save(sys.argv[1])
sys.path.insert(0, str(Path(__file__).resolve().parent))
from montar_docx import sanear  # noqa: E402
sanear(Path(sys.argv[1]))
