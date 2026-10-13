"""Markdown limpo -> manuscrito .docx no formato da IJGSI.

Uso: python word/montar_docx.py limpo.md saida.docx
Requer pandoc, python-docx e word/apa.csl.

Formato (Author Guidelines da IJGSI): Times New Roman 12, espaço 1,5, margens de
2,54 cm, número de página no canto superior direito, título em negrito, itálico,
maiúsculas e justificado, resumo < 1000 caracteres, 5 a 7 palavras-chave em
itálico e espaço simples, referências APA 7 em página nova, espaço duplo e recuo
deslocado de 1,25 cm, tabelas como texto editável, propriedades do arquivo sem
autor.
"""
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
FONTE = "Times New Roman"


def fonte(estilo, negrito=None, italico=None, tamanho=12):
    f = estilo.font
    f.name = FONTE
    f.size = Pt(tamanho)
    f.color.rgb = RGBColor(0, 0, 0)
    if negrito is not None:
        f.bold = negrito
    if italico is not None:
        f.italic = italico
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)  # rFonts tem de ser o primeiro filho de rPr
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONTE)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(a), None)


def paragrafo(estilo, espaco=1.5, alinhamento=None, antes=0, depois=0, recuo_primeira=None, recuo_esq=None):
    pf = estilo.paragraph_format
    pf.line_spacing = espaco
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    if alinhamento is not None:
        pf.alignment = alinhamento
    if recuo_primeira is not None:
        pf.first_line_indent = recuo_primeira
    if recuo_esq is not None:
        pf.left_indent = recuo_esq


def referencia_doc(destino: Path) -> None:
    subprocess.run(["pandoc", "-o", str(destino), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(destino)

    class Estilos:
        """Busca estilos pelo nome exibido (o reference.docx do pandoc usa ids diferentes)."""
        def __init__(self, styles):
            self.s = styles
        def __getitem__(self, nome):
            for x in self.s:
                if x.name == nome:
                    return x
            raise KeyError(nome)
        def add_style(self, *a):
            return self.s.add_style(*a)
    st = Estilos(doc.styles)
    for nome in ["Normal", "Body Text", "First Paragraph", "Compact", "Block Text", "Footnote Text",
                 "Caption", "Table Caption", "Image Caption", "Figure", "Captioned Figure", "Bibliography",
                 "Abstract", "Title", "Subtitle", "Author", "Date", "Hyperlink"]:
        try:
            s = st[nome]
        except KeyError:
            continue
        if s.type == WD_STYLE_TYPE.CHARACTER:
            continue
        fonte(s)
    for nome in ["Body Text", "First Paragraph", "Normal"]:
        paragrafo(st[nome], 1.5, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 0, Cm(1.25))
    paragrafo(st["Compact"], 1.0, None, 0, 0)
    # títulos de seção no estilo APA: nível 1 centralizado e negrito, nível 2 à esquerda e negrito
    for nome, alin, ital in (("Heading 1", WD_ALIGN_PARAGRAPH.CENTER, False), ("Heading 2", WD_ALIGN_PARAGRAPH.LEFT, False),
                             ("Heading 3", WD_ALIGN_PARAGRAPH.LEFT, True)):
        s = st[nome]
        fonte(s, negrito=True, italico=ital)
        paragrafo(s, 1.5, alin, 12, 0, Cm(0))
        s.paragraph_format.keep_with_next = True
    # estilos próprios da primeira página
    def novo(nome, base="Normal"):
        try:
            return st[nome]
        except KeyError:
            s = st.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st[base]
            return s
    s = novo("Titulo IJGSI"); fonte(s, True, True); paragrafo(s, 1.5, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 12, Cm(0))
    s.font.all_caps = True
    s = novo("Rotulo resumo"); fonte(s, True, False); paragrafo(s, 1.5, WD_ALIGN_PARAGRAPH.LEFT, 6, 0, Cm(0))
    s = novo("Resumo IJGSI"); fonte(s); paragrafo(s, 1.5, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 6, Cm(0))
    s = novo("Palavras-chave IJGSI"); fonte(s, False, True); paragrafo(s, 1.0, WD_ALIGN_PARAGRAPH.JUSTIFY, 6, 12, Cm(0))
    doc.save(destino)


def campo_pagina(par):
    run = par.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
            run._r.append(el)
        elif texto == "PAGE":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
            run._r.append(el)
        else:
            el = OxmlElement("w:t")
            el.text = texto
            run._r.append(el)
    run.font.name = FONTE
    run.font.size = Pt(12)


def bordas_apa(tabela):
    """Tabela APA: linhas só no topo, abaixo do cabeçalho e no fim."""
    tbl = tabela._tbl
    tblpr = tbl.tblPr
    antigo = tblpr.find(qn("w:tblBorders"))
    if antigo is not None:
        tblpr.remove(antigo)
    b = OxmlElement("w:tblBorders")
    for lado, val in (("top", "single"), ("left", "nil"), ("bottom", "single"), ("right", "nil"),
                      ("insideH", "nil"), ("insideV", "nil")):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), val)
        if val == "single":
            e.set(qn("w:sz"), "8")
            e.set(qn("w:color"), "000000")
        b.append(e)
    tblpr.append(b)
    cab = tabela.rows[0]
    for cel in cab.cells:
        tcpr = cel._tc.get_or_add_tcPr()
        tb = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000")
        tb.append(e)
        tcpr.append(tb)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER


# ordem dos filhos exigida pelo esquema OOXML (ECMA-376) nos elementos que editamos
ORDEM = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
           "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
           "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
           "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs "
           "em lang eastAsianLayout specVanish oMath",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "style": "name aliases basedOn next link autoRedefine hidden uiPriority semiHidden unhideWhenUsed qFormat locked "
             "personal personalCompose personalReply rsid pPr rPr tblPr trPr tcPr tblStylePr",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "settings": "writeProtection view zoom removePersonalInformation removeDateAndTime doNotDisplayPageBoundaries "
                "displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth printFormsData "
                "embedTrueTypeFonts embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins alignBordersAndEdges "
                "bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop hideSpellingErrors "
                "hideGrammaticalErrors activeWritingStyle proofState formsDesign attachedTemplate linkStyles "
                "stylePaneFormatFilter stylePaneSortMethod documentType mailMerge revisionView trackRevisions "
                "doNotTrackMoves doNotTrackFormatting documentProtection autoFormatOverride styleLockTheme "
                "styleLockQFSet defaultTabStop autoHyphenation consecutiveHyphenLimit hyphenationZone "
                "doNotHyphenateCaps showEnvelope summaryLength clickAndTypeStyle defaultTableStyle evenAndOddHeaders "
                "bookFoldRevPrinting bookFoldPrinting bookFoldPrintingSheets drawingGridHorizontalSpacing "
                "drawingGridVerticalSpacing displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery "
                "doNotUseMarginsForDrawingGridOrigin drawingGridHorizontalOrigin drawingGridVerticalOrigin "
                "doNotShadeFormData noPunctuationKerning characterSpacingControl printTwoOnOne strictFirstAndLastChars "
                "noLineBreaksAfter noLineBreaksBefore savePreviewPicture doNotValidateAgainstSchema "
                "saveInvalidXml ignoreMixedContent alwaysShowPlaceholderText doNotDemarcateInvalidXml "
                "saveXmlDataOnly useXSLTWhenSaving saveThroughXslt showXMLTags alwaysMergeEmptyNamespace "
                "updateFields hdrShapeDefaults footnotePr endnotePr compat docVars rsids mathPr attachedSchema "
                "themeFontLang clrSchemeMapping doNotIncludeSubdocsInStats doNotAutoCompressPictures forceUpgrade "
                "captions readModeInkLockDown smartTagType schemaLibrary shapeDefaults doNotEmbedSmartTags "
                "decimalSymbol listSeparator",
}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
# equações (OMML): ordem dos filhos nas propriedades de delimitador e de n-ário
ORDEM_M = {"dPr": "begChr sepChr endChr grow shp ctrlPr",
           "naryPr": "chr limLoc grow subHide supHide ctrlPr"}


def reordenar(raiz) -> None:
    for ns, tabela in ((W, ORDEM), (M, ORDEM_M)):
        _reordenar(raiz, ns, tabela)


def _reordenar(raiz, ns, tabela) -> None:
    for nome, ordem in tabela.items():
        pos = {n: i for i, n in enumerate(ordem.split())}
        for el in raiz.iter(ns + nome):
            filhos = list(el)
            filhos.sort(key=lambda c: pos.get(c.tag.split("}")[-1], len(pos)))
            for c in filhos:
                el.remove(c)
                el.append(c)


def sanear(docx: Path) -> None:
    """Corrige o XML para o esquema OOXML: ordem dos elementos, texto solto no
    styles.xml do modelo do pandoc, nsid curto na numeração e atributos de pgMar."""
    import re
    import zipfile
    from lxml import etree
    tmp = docx.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            if item.filename in ("word/document.xml", "word/styles.xml", "word/settings.xml", "word/numbering.xml"):
                raiz = etree.fromstring(dados)
                for el in raiz.iter():
                    if el.tag in (W + "rPr", W + "pPr", W + "tblPr", W + "tcPr"):
                        if el.text and el.text.strip():
                            el.text = None
                        for c in el:
                            if c.tail and c.tail.strip():
                                c.tail = None
                reordenar(raiz)
                for n in raiz.iter(W + "nsid"):
                    v = n.get(W + "val")
                    if v and len(v) < 8:
                        n.set(W + "val", v.zfill(8))
                for z in raiz.iter(W + "zoom"):
                    if z.get(W + "percent") is None:
                        z.set(W + "percent", "100")
                for m in raiz.iter(W + "pgMar"):
                    for a in ("header", "footer", "gutter"):
                        if m.get(W + a) is None:
                            m.set(W + a, "0" if a == "gutter" else "709")
                dados = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, dados)
    tmp.replace(docx)


def main() -> None:
    md, saida = Path(sys.argv[1]), Path(sys.argv[2])
    with tempfile.TemporaryDirectory() as tmp:
        ref = Path(tmp) / "ref.docx"
        referencia_doc(ref)
        subprocess.run(["pandoc", str(md), "--citeproc", "--bibliography", str(AQUI.parent / "referencias.bib"),
                        "--csl", str(AQUI / "apa.csl"), "--reference-doc", str(ref),
                        "--resource-path", f"{AQUI.parent}:{AQUI.parent / 'figuras'}", "-o", str(saida)], check=True)
    doc = Document(saida)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(2.54))
    sec.header_distance = Cm(1.25)
    sec.footer_distance = Cm(1.25)
    cab = sec.header.paragraphs[0]
    cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    campo_pagina(cab)

    # referências em página nova, espaço duplo, recuo deslocado de 1,25 cm
    em_refs = False
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading 1") and p.text.strip() == "References":
            p.paragraph_format.page_break_before = True
            em_refs = True
            continue
        if em_refs and p.text.strip():
            pf = p.paragraph_format
            pf.line_spacing = 2.0
            pf.left_indent = Cm(1.25)
            pf.first_line_indent = Cm(-1.25)
            pf.space_after = Pt(0)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # legendas APA (número e título acima) sem recuo e coladas no que vem depois
        txt = p.text.strip()
        if (txt.startswith("Table ") or txt.startswith("Figure ")) and len(txt) <= 9 and p.runs and p.runs[0].bold:
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.keep_with_next = True
            prox = p._p.getnext()
            if prox is not None:
                from docx.text.paragraph import Paragraph
                pp = Paragraph(prox, p._parent)
                pp.paragraph_format.first_line_indent = Cm(0)
                pp.paragraph_format.keep_with_next = True
                pp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(12)

    for t in doc.tables:
        bordas_apa(t)
        for linha in t.rows:
            for cel in linha.cells:
                for p in cel.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.name = FONTE
                        r.font.size = Pt(12)

    # propriedades do arquivo sem identificação
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.title = ""
    cp.subject = ""
    cp.keywords = ""
    cp.comments = ""
    doc.save(saida)
    sanear(saida)


if __name__ == "__main__":
    main()
