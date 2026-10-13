"""Material suplementar das revisões PRISMA (um .docx por artigo, anônimo).

Uso: python analise/suplemento_revisao.py
Gera artigos/<artigo>/word/submissao/supplementary-S1.docx e devolve, para o
Word completo, o Markdown das tabelas (função tabela_incluidos).
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "analise"))
sys.path.insert(0, str(RAIZ / "artigos" / "ijgsi-atratores" / "word"))
from revisao_prisma import REVISOES  # noqa: E402

ARTIGO = {"atratores": "ijgsi-atratores", "tabuleiro": "ijgsi-tabuleiro"}


def esc(t):
    return str(t or "").replace("|", "/").replace("*", "").strip()


def tabela_incluidos(rid: str) -> str:
    lista = json.loads((RAIZ / "artigos" / ARTIGO[rid] / "revisao_incluidos.json").read_text(encoding="utf8"))
    linhas = ["| Year | Title | Venue | DOI / arXiv | Category |", "|---|---|---|---|---|"]
    for r in sorted(lista, key=lambda r: (r["categoria"], -(r["ano"] or 0), r["titulo"])):
        ident = r["doi"] or (f"arXiv:{r['arxiv']}" if r["arxiv"] else "–")
        linhas.append(f"| {r['ano']} | {esc(r['titulo'])} | {esc(r['veiculo']) or '–'} | {ident} | {esc(r['categoria'])} |")
    return "\n".join(linhas)


def criterios(rid: str) -> str:
    cfg = REVISOES[rid]
    incl = "\n".join(f"- {x}" for x in cfg["inclusao"])
    excl = "\n".join(f"- {k}: {v}" for k, v in cfg["exclusao"].items())
    return f"**Inclusion criteria**\n\n{incl}\n\n**Exclusion reasons**\n\n{excl}\n"


def strings(rid: str) -> str:
    cfg = REVISOES[rid]
    b = json.loads((RAIZ / "experimentos" / "revisao" / rid / "busca.json").read_text(encoding="utf8"))
    c = json.loads((RAIZ / "experimentos" / "revisao" / rid / "contagens.json").read_text(encoding="utf8"))
    linhas = ["| Database | Field | Search string | Records |", "|---|---|---|---|"]
    for (nome, campo, chave) in (("OpenAlex", "title and abstract", "openalex"), ("arXiv", "abstract", "arxiv"),
                                 ("Semantic Scholar", "bulk search", "s2")):
        linhas.append(f"| {nome} | {campo} | `{esc(cfg[chave])}` | {c['porFonte'].get(nome, 0)} |")
    return (f"Search date: {b['data_da_busca']}. Period: {cfg['desde']} to {cfg['ate']}.\n\n" + "\n".join(linhas) + "\n")


def documento(rid: str) -> str:
    cfg = REVISOES[rid]
    c = json.loads((RAIZ / "experimentos" / "revisao" / rid / "contagens.json").read_text(encoding="utf8"))
    return f"""# Supplementary material S1. Systematic literature review (PRISMA 2020)

**Review question.** {cfg['pergunta']}

## Table S1. Databases and search strings

{strings(rid)}

## Eligibility criteria

{criterios(rid)}

## Screening process

Duplicates were removed by DOI, arXiv identifier and title similarity ({c['duplicatas']} records), and records without an abstract ({c['semResumo']}) could not be screened. Titles and abstracts of {c['triados']} records were screened by a language model (claude-sonnet-5) against the criteria above; {c['avaliados']} records were then assessed for eligibility with stricter instructions by a second model (claude-opus-5-5). One of the authors checked the included records as a second reviewer, which changed {c['decisoesManuais']} decision(s). Every decision, with its reason, is available in the project repository (link removed for anonymous review).

## Table S2. Studies included in the review (n = {c['incluidos']})

{tabela_incluidos(rid)}
"""


def main():
    from montar_docx import sanear
    from docx import Document
    from docx.shared import Cm, Pt
    for rid, art in ARTIGO.items():
        destino = RAIZ / "artigos" / art / "word" / "submissao" / "supplementary-S1.docx"
        with tempfile.TemporaryDirectory() as t:
            md = Path(t) / "s1.md"
            md.write_text(documento(rid), encoding="utf8")
            subprocess.run(["pandoc", str(md), "-o", str(destino)], check=True)
        doc = Document(destino)
        for s in doc.sections:
            s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)
        for st in doc.styles:
            if st.type == 1:
                st.font.name = "Times New Roman"
        for tb in doc.tables:
            for linha in tb.rows:
                for cel in linha.cells:
                    for p in cel.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9)
        cp = doc.core_properties
        cp.author = cp.last_modified_by = ""
        doc.save(destino)
        sanear(destino)
        print(destino)


if __name__ == "__main__":
    main()
