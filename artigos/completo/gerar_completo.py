"""Gera um único .docx com os dois artigos completos (versão de leitura, com autoria).

Uso: python artigos/completo/gerar_completo.py
Saída: artigos/completo/artigos-completos.docx

Cada artigo é convertido do seu main.tex pelo mesmo caminho da versão de
submissão (word/preparar_md.py), as citações são resolvidas separadamente por
artigo (cada um com a sua lista de referências APA 7) e o conjunto é montado com
o modelo de estilos de word/montar_docx.py: Times New Roman 12, espaço 1,5,
figuras e tabelas no estilo APA, número de página no canto superior direito.
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

AQUI = Path(__file__).resolve().parent
ARTIGOS = AQUI.parent
sys.path.insert(0, str(ARTIGOS / "ijgsi-atratores" / "word"))
from montar_docx import bordas_apa, campo_pagina, referencia_doc, sanear  # noqa: E402

REPO = "https://github.com/caiosilvaazeredo/cronicas-de-aethelgard"
ORDEM = [("ijgsi-atratores", "Article 1"), ("ijgsi-tabuleiro", "Article 2")]
QUEBRA = '\n\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'


def artigo_em_markdown(pasta: Path, tmp: Path) -> str:
    bruto, limpo, citado = tmp / f"{pasta.name}-bruto.md", tmp / f"{pasta.name}-limpo.md", tmp / f"{pasta.name}-citado.md"
    subprocess.run(["pandoc", str(pasta / "main.tex"), "-t", "markdown-simple_tables-multiline_tables-grid_tables+pipe_tables",
                    "--wrap=none", "-o", str(bruto)], check=True, cwd=pasta, capture_output=True)
    subprocess.run([sys.executable, str(pasta / "word" / "preparar_md.py"), str(bruto), str(limpo)], check=True, cwd=pasta,
                   capture_output=True)
    subprocess.run(["pandoc", str(limpo), "--citeproc", "--bibliography", str(pasta / "referencias.bib"),
                    "--csl", str(pasta / "word" / "apa.csl"), "-t", "markdown-citations", "--wrap=none", "-o", str(citado)],
                   check=True, cwd=pasta)
    md = citado.read_text(encoding="utf8")
    md = re.sub(r"^---\n.*?\n---\n", "", md, flags=re.S)  # metadados de cada artigo
    md = md.replace(r"\[repository link removed for anonymous review\]", REPO)
    # a lista de referências vira um bloco com o estilo Bibliography, sem as âncoras internas
    md = re.sub(r"^:+ \{#refs[^}]*\}\s*$", '::::: {custom-style="Bibliography"}', md, flags=re.M)
    md = re.sub(r"^:+ \{#ref-[^}]*\}\s*$\n?", "", md, flags=re.M)
    linhas, saida, aberto = md.split("\n"), [], False
    for l in linhas:
        if l.startswith('::::: {custom-style="Bibliography"}'):
            aberto = True
            saida.append(l)
            continue
        if aberto and re.fullmatch(r":{3,}\s*", l):
            continue  # fechamentos das entradas removidas
        saida.append(l)
    if aberto:
        saida.append(":::::")
    md = "\n".join(saida)
    md = re.sub(r"\]\(#ref-[^)]*\)", "]", md)  # links internos das citações
    return md


def capa() -> str:
    return """::: {custom-style="Titulo IJGSI"}
NARRATIVE ATTRACTORS AND BOARD-BASED WORLD MODELS: TWO ARTICLES ON STORY CONVERGENCE IN SHARED-WORLD ROLE-PLAY
:::

Caio Silva Azeredo and Geraldo Xexéo

Systems Engineering and Computer Science Program (PESC/COPPE), Federal University of Rio de Janeiro (UFRJ), Brazil

**Article 1.** Narrative attractors: why stories converge when language-model agents share a role-playing world.

**Article 2.** From board to graph: a hybrid human-AI role-playing board for critical media and AI literacy.

Both articles were prepared for *The International Journal of Games and Social Impact* (Vol. 5, No. 4, "Games, Gamification, Media Literacy and Transformation in Educational Contexts"). This compilation is a reading version with authorship; the anonymised submission files are generated separately. Each article keeps its own figures, tables and reference list (APA 7). Every reference was checked against its source: entries with a DOI against the Crossref record, the others against the publisher or repository page. The literature reviews follow PRISMA 2020; their search strings, records and screening decisions are in the repository (""" + REPO + """, folder `experimentos/revisao`).
"""


def main() -> None:
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        partes = [capa()]
        for nome, rotulo in ORDEM:
            partes.append(QUEBRA + artigo_em_markdown(ARTIGOS / nome, tmp))
        md = tmp / "completo.md"
        md.write_text("\n".join(partes), encoding="utf8")
        ref = tmp / "ref.docx"
        referencia_doc(ref)
        saida = AQUI / "artigos-completos.docx"
        caminhos = ":".join(str(ARTIGOS / n / "figuras") for n, _ in ORDEM)
        subprocess.run(["pandoc", str(md), "--reference-doc", str(ref), "--resource-path", caminhos, "-o", str(saida)],
                       check=True)
    doc = Document(saida)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(2.54))
    sec.header_distance = sec.footer_distance = Cm(1.25)
    cab = sec.header.paragraphs[0]
    cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    campo_pagina(cab)
    from docx.text.paragraph import Paragraph
    for p in doc.paragraphs:
        txt = p.text.strip()
        if p.style.name == "Bibliography":
            pf = p.paragraph_format
            pf.line_spacing, pf.left_indent, pf.first_line_indent = 1.15, Cm(1.25), Cm(-1.25)
            pf.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if p.style.name.startswith("Heading 1") and txt == "References":
            p.paragraph_format.page_break_before = True
        if (txt.startswith("Table ") or txt.startswith("Figure ")) and len(txt) <= 9 and p.runs and p.runs[0].bold:
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.keep_with_next = True
            prox = p._p.getnext()
            if prox is not None:
                pp = Paragraph(prox, p._parent)
                pp.paragraph_format.first_line_indent = Cm(0)
                pp.paragraph_format.keep_with_next = True
                pp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.first_line_indent = Cm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(12)
    for tb in doc.tables:
        bordas_apa(tb)
        for linha in tb.rows:
            for cel in linha.cells:
                for p in cel.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(11)
    cp = doc.core_properties
    cp.author = "Caio Silva Azeredo; Geraldo Xexéo"
    cp.title = "Narrative attractors and board-based world models"
    doc.save(saida)
    sanear(saida)
    print(saida)


if __name__ == "__main__":
    main()
